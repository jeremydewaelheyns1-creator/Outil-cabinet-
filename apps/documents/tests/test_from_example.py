"""Un modèle tiré d'un document déjà rempli (D-124). Tout est fictif."""

from __future__ import annotations

import io
from datetime import date

import pytest
from django.urls import reverse
from docx import Document as Word

from apps.clients.models import Relation
from apps.core.testing.factories import ClientFactory, PersonFactory
from apps.documents import from_example, letter
from apps.documents.models import LetterTemplate

pytestmark = pytest.mark.django_db


def filled_pv() -> bytes:
    d = Word()
    d.add_paragraph("Menuiserie Fictive SRL")
    d.add_paragraph("Rue des Essais 12")
    d.add_paragraph("5000 Namur")
    d.add_paragraph("BCE : BE 0987.654.394")
    d.add_paragraph("Procès-verbal de l'assemblée générale ordinaire du 30/06/2026")
    d.add_paragraph("La séance est ouverte à 18h00 sous la présidence de Paul Testeur, gérant.")
    d.add_paragraph("Associés présents : Paul Testeur (60 parts) et Marie Essai (40 parts).")
    d.add_paragraph("L'assemblée approuve les comptes annuels arrêtés au 31/12/2025, bénéfice de "
                    "32.174,60 €.")  # fmt: skip
    d.add_paragraph("Le bénéfice est reporté. Contact : paul@menuiserie-fictive.example")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Impôts", "8.930,00"
    d.add_paragraph("La séance est levée à 18h30.")
    d.add_paragraph("Fait à Namur, le 30 juin 2026.")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


@pytest.fixture
def example_client():
    from apps.closing import services

    company = ClientFactory(name="Menuiserie Fictive", legal_form="SRL",
                            enterprise_number="0987654394", street="Rue des Essais 12",
                            postal_code="5000", city="Namur", fiscal_year_end_month=12,
                            fiscal_year_end_day=31)  # fmt: skip
    manager = PersonFactory(name="Testeur", first_name="Paul")
    partner = PersonFactory(name="Essai", first_name="Marie")
    Relation.objects.create(entity=company, person=manager, role=Relation.Role.MANAGER)
    Relation.objects.create(entity=company, person=manager, role=Relation.Role.SHAREHOLDER)
    Relation.objects.create(entity=company, person=partner, role=Relation.Role.SHAREHOLDER)
    services.ensure_years(company, date(2026, 10, 30))
    return company


def test_everything_of_the_client_becomes_a_hole(example_client):
    result = from_example.convert(filled_pv(), example_client)
    text = from_example.text_of(result.content)
    for leftover in ("Menuiserie Fictive", "Essais 12", "Namur", "0987", "Paul Testeur",
                     "Marie Essai", "30/06/2026", "18h00", "18h30", "31/12/2025", "32.174,60",
                     "8.930,00", "paul@", "30 juin 2026"):  # fmt: skip
        assert leftover not in text, leftover
    assert "Procès-verbal de l'assemblée générale ordinaire" in text  # le texte du cabinet reste
    report = dict(result.report())
    assert report["nom du client"] >= 1 and report["montant"] == 2 and report["heure"] == 2
    holes = from_example.holes_for(result.content, "pv_ag")
    fields = [v["field"] for v in holes.values()]
    for code in ("client.denomination", "client.adresse", "client.localite", "client.numero",
                 "client.representant", "exercice.cloture", "ag.date", "ag.debut", "ag.fin",
                 "demander"):  # fmt: skip
        assert code in fields, code


def test_upload_a_filled_document_and_reuse_it_for_another_client(client, make_user,
                                                                   example_client):  # fmt: skip
    from django.core.files.uploadedfile import SimpleUploadedFile

    admin = make_user(is_staff=True, capabilities=["clients.edit"])
    client.force_login(admin)
    upload = SimpleUploadedFile("pv-menuiserie.docx", filled_pv())
    # Aucun client à choisir : retrouvé par le numéro d'entreprise imprimé.
    response = client.post(reverse("lettre:models"), {"sorte": "pv_ag", "modele": upload})
    template = LetterTemplate.objects.get()
    assert response.status_code == 302 and template.kind == "pv_ag" and not template.active
    assert template.origin["exemple"] and "Menuiserie" not in str(template.origin)
    page = client.get(reverse("lettre:model", args=[template.pk])).content.decode()
    assert "Modèle tiré d'un document rempli" in page
    listing = client.get(reverse("lettre:models") + "?sorte=pv_ag").content.decode()
    assert "vide ou déjà rempli" in listing and "Fait pour le client" not in listing
    assert client.get(reverse("lettre:download", args=[template.pk])).status_code == 200
    # Réutilisé pour un autre client : son nom apparaît, plus rien de la Menuiserie.
    other = ClientFactory(name="Garage Autre", legal_form="SRL", street="Avenue B 3",
                          postal_code="4000", city="Liège")  # fmt: skip
    answers = {k: "1,00" for k, _ in letter.questions(template)}
    ctx = {"date": "15/06/2027", "debut": "17h00", "fin": "17h45", "exercice_fin": "31/12/2026",
           "dividende": "non"}  # fmt: skip
    text = from_example.text_of(letter.fill(template, other, answers, ctx))
    assert "Garage Autre" in text and "Avenue B 3" in text and "17h45" in text
    assert "Menuiserie" not in text and "Namur" not in text and "Testeur" not in text


def test_without_knowing_the_client_the_shape_is_enough():
    """Client inconnu de l'outil : la forme suffit (société, numéro, localité, dates, heures,
    montants) ; l'assistant IA trouve les noms de personnes."""
    result = from_example.convert(filled_pv(), None)
    text = from_example.text_of(result.content)
    for leftover in ("Menuiserie Fictive SRL", "0987", "5000 Namur", "30/06/2026", "18h00",
                     "32.174,60", "paul@"):  # fmt: skip
        assert leftover not in text, leftover
    holes = from_example.holes_for(result.content, "pv_ag")
    assert "exercice.cloture" in [v["field"] for v in holes.values()]  # « arrêtés au … »
