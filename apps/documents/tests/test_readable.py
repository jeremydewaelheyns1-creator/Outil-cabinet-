"""Tous les formats lisibles, pas seulement le PDF (D-100) : XML d'Intervat, XBRL, Biztax,
autre XML, Excel, Word — dans le visualiseur et sur la page « Doublon à vérifier »."""

import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from apps.core.testing.factories import ClientFactory
from apps.core.testing.numbers import enterprise_number
from apps.documents import readable, services
from apps.documents.models import Document, InboxItem
from apps.documents.tests.test_official_files import accounts_xbrl, biztax, xbrl_upload

pytestmark = pytest.mark.django_db
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Grand livre"
    sheet.append(["Compte", "Libellé", "Montant"])
    sheet.append(["604000", "Achats fictifs", 1234.5])
    book.create_sheet("Vide")
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def docx() -> bytes:
    from docx import Document as WordDocument

    word = WordDocument()
    word.add_paragraph("Procès-verbal fictif de l'assemblée générale")
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Résolution"
    table.rows[0].cells[1].text = "Approuvée"
    out = io.BytesIO()
    word.save(out)
    return out.getvalue()


OTHER_XML = (b"<?xml version='1.0'?><Fiche xmlns='urn:fictif'><Nom>Dossier fictif</Nom>"
             b"<Annexe>" + b"A" * 500 + b"</Annexe></Fiche>")  # fmt: skip


def stored(client, name, content):
    return services.store(client, SimpleUploadedFile(name, content), title=name,
                          folder="administratif")  # fmt: skip


def viewer(web, document):
    url = reverse("clients:document_viewer", args=[document.client_id, document.pk])
    return web.get(url, HTTP_HX_REQUEST="true").content.decode()


@pytest.mark.parametrize(
    ("name", "content", "expected"),
    [
        ("comptes.xbrl", lambda: accounts_xbrl(enterprise_number(9101)), "Chiffres clés"),
        ("isoc.biztax", lambda: biztax(enterprise_number(9101)), "Résultat fiscal"),
        ("autre.xml", lambda: OTHER_XML, "Nom : Dossier fictif"),
        ("balance.xlsx", xlsx, "Achats fictifs"),
        ("pv.docx", docx, "Procès-verbal fictif"),
    ],
)
def test_every_format_is_readable_in_the_viewer(web, name, content, expected):
    company = ClientFactory(enterprise_number=enterprise_number(9101))
    html = viewer(web, stored(company, name, content()))
    assert expected in html.replace("&#x27;", "'")
    assert "Pas d'aperçu" not in html


def test_details():
    company = ClientFactory()
    tree = readable.of(stored(company, "autre.xml", OTHER_XML))
    assert tree is not None and tree.kind == "xml"
    assert any("(500 caractères)" in line for line in tree.lines)  # valeur longue raccourcie
    sheets = readable.of(stored(company, "b.xlsx", xlsx()))
    assert sheets is not None and [s.name for s in sheets.sheets] == ["Grand livre", "Vide"]
    assert sheets.sheets[0].rows[1] == ["604000", "Achats fictifs", "1234.5"]
    word = readable.of(stored(company, "pv.docx", docx()))
    assert word is not None and "Résolution | Approuvée" in word.lines


def test_broken_or_dangerous_files_have_no_preview_but_no_error(web):
    company = ClientFactory()
    bomb = (b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;">]>'
            b"<x>&b;</x>")  # fmt: skip
    for name, content in (("bombe.xml", bomb), ("abime.xml", b"<?xml version='1.0'?><a><b>")):
        document = stored(company, name, content)
        assert readable.of(document) is None
        assert "Pas d'aperçu" in viewer(web, document)
    fake = Document.objects.create(client=company, title="Faux Excel", folder="administratif",
                                   original_name="f.xlsx", content_type=XLSX, size=10,
                                   sha256="0" * 64,
                                   file=SimpleUploadedFile("f.xlsx", b"PK\x03\x04x"))  # fmt: skip
    assert readable.of(fake) is None


@pytest.mark.parametrize(
    ("name", "content", "expected"),
    [
        ("isoc.biztax", lambda: biztax(enterprise_number(9102)), "Résultat fiscal"),
        ("autre.xml", lambda: OTHER_XML, "Nom : Dossier fictif"),
        ("balance.xlsx", xlsx, "Achats fictifs"),
    ],
)
def test_duplicate_page_shows_both_sides_for_any_format(web, name, content, expected):
    company = ClientFactory(enterprise_number=enterprise_number(9102))
    stored(company, name, content())
    web.post(reverse("depot:index"), {"files": [xbrl_upload(name, content())]})
    item = InboxItem.objects.get()
    assert item.duplicate_of is not None
    html = web.get(reverse("depot:duplicate", args=[item.pk])).content.decode()
    assert html.count(expected) == 2  # à gauche ET à droite
    assert "Pas d'aperçu" not in html
