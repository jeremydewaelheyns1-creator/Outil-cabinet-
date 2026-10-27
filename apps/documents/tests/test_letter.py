"""Lettre de mission sur le modèle Word du cabinet (D-069). Modèle fictif construit en code."""

import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client as Browser
from django.urls import reverse
from docx import Document as Word

from apps.core.models import AccessLog
from apps.documents import letter
from apps.documents.models import Document, LetterTemplate

from .letter_fixture import install, model
from .test_generation import company, text_of

pytestmark = pytest.mark.django_db


@pytest.fixture
def admin(make_user):
    browser = Browser()
    browser.force_login(make_user(is_staff=True, capabilities=["clients.edit"]))
    return browser


def upload(admin, content=None):
    file = SimpleUploadedFile("lettre.docx", content or model())
    response = admin.post(reverse("lettre:models"), {"modele": file})
    assert response.status_code == 302
    return LetterTemplate.objects.get()


def configure(admin, template, **extra):
    """Accepte les propositions de l'outil, complète le reste et coche les vérifications."""
    data = {
        "action": "activer",
        "verifie": [c.code for c in letter.checks(bytes(template.content), template.holes)],
    }
    for hole in letter.holes(bytes(template.content)):
        data[f"trou-{hole.key}"] = (template.holes.get(hole.key) or {}).get("field") or letter.ASK
        data[f"libelle-{hole.key}"] = f"Question {hole.key}"
    data.update(extra)
    return admin.post(reverse("lettre:model", args=[template.pk]), data)


def test_holes_are_found_and_guessed_from_the_text_before():
    found = {h.key: h.guess for h in letter.holes(model())}
    assert list(found.values())[:4] == [
        "client.denomination",
        "client.siege",
        "client.numero",
        "client.representant",
    ]
    assert letter.ASK in found.values()  # honoraires, lieu
    questions = {h.question for h in letter.holes(model()) if h.guess == letter.ASK}
    assert {
        "Honoraires (montant hors TVA)",
        "Lieu de signature",
        "Premier exercice concerné",
    } <= questions
    texts = [h.text for h in letter.holes(model())]
    assert "[année]" in texts and "à adapter" in texts  # crochets et surlignage


def test_only_admins_manage_the_model(web):
    assert web.get(reverse("lettre:models")).status_code == 403


def test_missing_mandatory_element_blocks_activation(admin):
    template = upload(admin, model(complete=False))
    configure(admin, template)
    template.refresh_from_db()
    assert not template.active
    assert any("résili" in p.lower() or "fin" in p.lower() for p in letter.problems(template))


def test_unchecked_element_blocks_then_activation(admin):
    template = upload(admin)
    configure(admin, template, verifie=[])
    template.refresh_from_db()
    assert not template.active
    configure(admin, template)
    template.refresh_from_db()
    assert template.active
    assert AccessLog.objects.filter(action=AccessLog.Action.IMPORT).exists()


def test_letter_is_filled_for_the_client_and_options_follow_the_dossier(admin):
    template = upload(admin)
    tva = next(i for i, t in letter.paragraphs(bytes(template.content)) if "TVA périodiques" in t)
    configure(admin, template, **{f"option-{tva}": "personne"})
    template.refresh_from_db()
    assert template.active
    client, _ = company()
    questions = dict(letter.questions(template))
    answers = {key: "1 200" if "4:" in key else "Namur" for key in questions}
    word = Word(io.BytesIO(letter.fill(template, client, answers)))
    text = "\n".join(p.text for p in word.paragraphs)
    assert "Société Fictive" in text and "Anne Fictif" in text and "1 200 EUR" in text
    assert "XXX" not in text
    assert "TVA périodiques" not in text  # réservé aux personnes physiques


def test_prepare_page_generates_the_draft_from_the_cabinet_model(admin, web):
    template = upload(admin)
    configure(admin, template)
    client, _ = company()
    url = reverse("clients:letter_prepare", args=[client.pk])
    page = web.get(url)
    assert page.status_code == 200 and "Question 4:0" in page.content.decode()
    keys = [key for key, _ in letter.questions(LetterTemplate.objects.get())]
    assert web.post(url, {}).status_code == 200  # réponses manquantes
    response = web.post(url, dict.fromkeys(keys, "Valeur fictive"))
    document = Document.objects.get(client=client, doc_type="lettre_mission")
    assert response.status_code == 302 and document.state == Document.State.DRAFT
    assert "Société Fictive" in text_of(document)


def test_direct_generate_url_refuses_the_letter(web):
    install()
    client, _ = company()
    url = reverse("clients:document_generate", args=[client.pk, "lettre_mission"])
    assert web.post(url).status_code == 404  # passe par « Préparer »


def test_not_a_word_file_is_refused(admin):
    file = SimpleUploadedFile("lettre.docx", b"pas un word")
    admin.post(reverse("lettre:models"), {"modele": file})
    assert not LetterTemplate.objects.exists()


def test_zip_bomb_is_refused_before_opening(admin):
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<x/>")
        archive.writestr("word/bomb.xml", b"0" * (60 * 1024 * 1024))  # 60 Mo, compressé en Ko
    content = buffer.getvalue()
    assert len(content) < 200_000
    file = SimpleUploadedFile("lettre.docx", content)
    response = admin.post(reverse("lettre:models"), {"modele": file}, follow=True)
    assert "anormal" in response.content.decode() and not LetterTemplate.objects.exists()


def test_split_holes_and_hyperlinks_are_filled_in_place():
    """Word coupe souvent « XXX » en plusieurs morceaux de mise en forme ; un lien dans le
    paragraphe ne doit pas décaler le remplissage."""
    from docx.oxml import OxmlElement

    from apps.core.testing.factories import ClientFactory

    doc = Word()
    p = doc.add_paragraph("Et la société ")
    p.add_run("XX").bold = True
    p.add_run("X, dont le siège social est situé ")
    p.add_run("XXX").italic = True
    p.add_run(".")
    p2 = doc.add_paragraph("Voir ")
    link, run, text = OxmlElement("w:hyperlink"), OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "notre site"
    run.append(text)
    link.append(run)
    p2._p.append(link)
    p2.add_run(" ; honoraires : XXX EUR.")
    out = io.BytesIO()
    doc.save(out)
    content = out.getvalue()
    found = letter.holes(content)
    template = LetterTemplate(content=content, options={},
                              holes={h.key: {"field": h.guess} for h in found})  # fmt: skip
    client = ClientFactory(name="Fictive Alpha", street="Rue du Test", house_number="1",
                           postal_code="5000", city="Namur")  # fmt: skip
    filled = Word(io.BytesIO(letter.fill(template, client, {"trou:1:0": "900"})))
    texts = [p.text for p in filled.paragraphs]
    assert texts[0].startswith(f"Et la société {client.display_name}, dont le siège social")
    assert "Rue du Test 1, 5000 Namur." in texts[0] and "XX" not in texts[0]
    assert texts[1] == "Voir notre site ; honoraires : 900 EUR."


def test_option_alone_in_a_table_cell_keeps_the_word_file_valid():
    from apps.core.testing.factories import ClientFactory

    doc = Word()
    doc.add_paragraph("Lettre de mission pour XXX.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).paragraphs[0].text = "Mission TVA : déclarations périodiques."
    table.cell(0, 1).paragraphs[0].text = "Autre case."
    out = io.BytesIO()
    doc.save(out)
    content = out.getvalue()
    index = next(i for i, t in letter.paragraphs(content) if t.startswith("Mission TVA"))
    template = LetterTemplate(content=content, holes={"0:0": {"field": "client.denomination"}},
                              options={str(index): {"condition": "tva", "end": index}})  # fmt: skip
    client = ClientFactory(vat_status="not_subject")
    filled = Word(io.BytesIO(letter.fill(template, client, {})))
    cell = filled.tables[0].cell(0, 0)
    assert len(cell.paragraphs) == 1 and cell.paragraphs[0].text == ""
    assert filled.tables[0].cell(0, 1).paragraphs[0].text == "Autre case."
