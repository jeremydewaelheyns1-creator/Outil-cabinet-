"""Modèle Word fictif de lettre de mission, construit en code (jamais un vrai modèle)."""

from __future__ import annotations

import io

from docx import Document as Word
from docx.enum.text import WD_COLOR_INDEX


def model(complete: bool = True) -> bytes:
    doc = Word()
    doc.add_heading("Lettre de mission", level=1)
    doc.add_paragraph("Entre le cabinet Exemple, représenté par son associé, d'une part,")
    doc.add_paragraph(
        "Et la société XXX, dont le siège social est situé XXX, inscrite à la "
        "Banque-Carrefour des Entreprises sous le numéro XXX, représentée par XXX, "
        "client, d'autre part."
    )
    doc.add_paragraph("La présente mission prend cours à partir de l'exercice [année].")
    doc.add_paragraph("Les honoraires sont fixés à XXX EUR hors TVA, payables à 30 jours.")
    doc.add_paragraph("Le cabinet établit les déclarations TVA périodiques du client.")
    p = doc.add_paragraph("Une remarque libre : ")
    run = p.add_run("à adapter")
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph("Le client remet ses documents ; le cabinet les conserve 7 ans.")
    if complete:
        doc.add_paragraph(
            "Chaque partie peut résilier la convention moyennant un préavis de 3 mois."
        )
        doc.add_paragraph("Tout litige relève des tribunaux de l'arrondissement du cabinet.")
    doc.add_paragraph("Le cabinet respecte la loi anti-blanchiment et le secret professionnel.")
    doc.add_paragraph("Fait à XXX, le XXX, en deux exemplaires. Lu et approuvé, signature.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def install(complete: bool = True):  # type: ignore[no-untyped-def]
    """Dépose et active le modèle fictif, comme un administrateur du cabinet."""
    from apps.documents import letter
    from apps.documents.models import LetterTemplate

    content = model(complete)
    holes = {h.key: {"field": h.guess or letter.KEEP, "label": ""} for h in letter.holes(content)}
    checks = [c.code for c in letter.checks(content, holes) if c.excerpt]
    return LetterTemplate.objects.create(
        name="modele-fictif.docx", content=content, holes=holes, checks=checks, active=True
    )
