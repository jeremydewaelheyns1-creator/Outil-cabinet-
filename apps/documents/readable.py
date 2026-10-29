"""Aperçu lisible des fichiers qui ne sont pas des images de pages (D-097, D-098, D-100).

Le visualiseur montre les PDF et les photos page par page (`preview`). Pour les autres
formats, on montre le contenu en clair, sans jamais exécuter quoi que ce soit :

- XML d'Intervat (déclaration TVA, listing IC) → grilles en tableau ;
- XBRL de la Banque Nationale (comptes annuels) et Biztax → chiffres clés ;
- autre XML → l'arbre des balises, indenté (valeurs longues raccourcies) ;
- Excel (.xlsx) → les premières lignes de chaque feuille ;
- Word (.docx) → le texte et les tableaux.

Fonctionne pour un document rangé comme pour un fichier encore à classer (doublons).
Lecture sûre : defusedxml pour le XML ; taille décompressée bornée pour Word et Excel.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from typing import Any

from defusedxml import ElementTree as SafeET

from .models import Document, InboxItem

MAX_BYTES = 15_000_000  # au-delà : pas d'aperçu, on télécharge
MAX_UNZIPPED = 60_000_000  # Word / Excel : taille une fois décompressés
XML_LINES = 400
VALUE_CHARS = 120
SHEET_ROWS, SHEET_COLS, SHEETS = 60, 15, 5
DOC_BLOCKS = 300
XML_TYPES = ("application/xml", "text/xml")
XSI = "{http://www.w3.org/2001/XMLSchema-instance}"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass
class Sheet:
    name: str
    rows: list[list[str]]
    more: bool = False


@dataclass
class Readable:
    kind: str  # « intervat », « xbrl », « xml », « sheets », « text »
    data: Any = None
    lines: list[str] = field(default_factory=list)
    sheets: list[Sheet] = field(default_factory=list)
    truncated: bool = False


def _local(tag: object) -> str:
    return str(tag).rsplit("}", 1)[-1]


def _xml_tree(content: bytes) -> Readable | None:
    try:
        root = SafeET.fromstring(content)
    except Exception:  # fichier abîmé ou dangereux : pas d'aperçu
        return None
    lines: list[str] = []

    def walk(node: Any, depth: int) -> None:
        if len(lines) >= XML_LINES:
            return
        text = (node.text or "").strip()
        if len(text) > VALUE_CHARS:
            text = f"{text[:VALUE_CHARS]}… ({len(text)} caractères)"
        attrs = " ".join(
            f'{_local(k)}="{v[:60]}"' for k, v in node.attrib.items() if not str(k).startswith(XSI)
        )
        head = _local(node.tag) + (f" [{attrs}]" if attrs else "")
        lines.append("  " * depth + head + (f" : {text}" if text else ""))
        for child in node:
            walk(child, depth + 1)

    walk(root, 0)
    return Readable("xml", lines=lines, truncated=len(lines) >= XML_LINES)


def _zip_ok(content: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            return sum(i.file_size for i in archive.infolist()) <= MAX_UNZIPPED
    except zipfile.BadZipFile:
        return False


def _cell(value: object) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return str(value.strftime("%d/%m/%Y"))
    return str(value)[:VALUE_CHARS]


def _sheets(content: bytes) -> Readable | None:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return None
    sheets = []
    try:
        for sheet in book.worksheets[:SHEETS]:
            rows: list[list[str]] = []
            more = False
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell(v) for v in row[:SHEET_COLS]]
                if not any(cells):
                    continue
                if len(rows) >= SHEET_ROWS:
                    more = True
                    break
                rows.append(cells)
            width = max((len(r) for r in rows), default=0)
            while width and all(not r[width - 1] for r in rows if len(r) >= width):
                width -= 1  # colonnes vides à droite
            sheets.append(Sheet(sheet.title, [(r + [""] * width)[:width] for r in rows], more))
    finally:
        book.close()
    return Readable("sheets", sheets=sheets, truncated=len(book.worksheets) > SHEETS)


def _text(content: bytes) -> Readable | None:
    from docx import Document as WordDocument

    try:
        word = WordDocument(io.BytesIO(content))
    except Exception:
        return None
    lines = [p.text for p in word.paragraphs if p.text.strip()]
    for table in word.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return Readable("text", lines=lines[:DOC_BLOCKS], truncated=len(lines) > DOC_BLOCKS)


def of(source: Document | InboxItem) -> Readable | None:
    """Aperçu lisible du fichier, ou None (PDF, images : vues page par page ailleurs)."""
    content_type = source.content_type
    if content_type not in (*XML_TYPES, XLSX, DOCX) or source.size > MAX_BYTES:
        return None
    with source.file.open("rb") as handle:
        content = handle.read()
    if content_type in XML_TYPES:
        from apps.vat import intervat

        from . import xbrl

        try:
            declaration = intervat.preview(content)
        except Exception:  # XML qui n'est pas d'Intervat
            declaration = None
        if declaration is not None:
            return Readable("intervat", data=declaration)
        official = xbrl.read(content)
        if official is not None:
            return Readable("xbrl", data=official)
        return _xml_tree(content)
    if not _zip_ok(content):
        return None
    return _sheets(content) if content_type == XLSX else _text(content)
