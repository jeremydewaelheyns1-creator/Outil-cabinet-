"""PV d'assemblée générale préparé depuis la case AG (D-120). Modèle Word fictif, construit en
code avec la même structure qu'un PV de cabinet (trous entre crochets)."""

from __future__ import annotations

import io
from datetime import date

import pytest
from django.urls import reverse

from apps.closing import ag, deposits, services
from apps.closing.models import ClosingYear
from apps.core.testing.factories import ClientFactory, PersonFactory

pytestmark = pytest.mark.django_db


def model_docx() -> bytes:
    from docx import Document as Word

    doc = Word()
    doc.add_paragraph("[SOCIÉTÉ] S.R.L.")
    doc.add_paragraph("[ADRESSE]")
    doc.add_paragraph("[LOCALITÉ]")
    doc.add_paragraph("R.P.M. : [IDENTIFIANT]")
    doc.add_paragraph("Procès-verbal de l'assemblée générale ordinaire des associés")
    doc.add_paragraph("L'an [ANNÉE] le [DATE], s'est tenue au siège social l'assemblée générale.")
    doc.add_paragraph("La séance est ouverte à [HEURE] sous la présidence de [PERSONNE 1].")
    doc.add_paragraph("Examen et approbation des comptes annuels arrêtés au [DATE DE CLÔTURE].")
    doc.add_paragraph("Bénéfice net de l'exercice à affecter : [MONTANT] €")
    doc.add_paragraph("L'assemblée décide de reporter l'intégralité du bénéfice.")
    doc.add_paragraph("L'assemblée décide de distribuer un dividende brut de [MONTANT] €.")
    doc.add_paragraph("L'AGO confirme le mandat donné à [CABINET COMPTABLE].")
    doc.add_paragraph("La séance est levée à [HEURE].")
    doc.add_paragraph("Fait à [LOCALITÉ] le [DATE]")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


@pytest.fixture
def admin_web(client, make_user):
    admin = make_user(is_staff=True, capabilities=["clients.edit"])
    client.force_login(admin)
    return client


def install(admin_web):
    from django.core.files.uploadedfile import SimpleUploadedFile

    from apps.documents import letter
    from apps.documents.models import LetterTemplate

    upload = SimpleUploadedFile("pv-fictif.docx", model_docx())
    admin_web.post(reverse("lettre:models"), {"sorte": "pv_ag", "modele": upload})
    template = LetterTemplate.objects.get(kind="pv_ag")
    fields = {h.key: template.holes.get(h.key, {}).get("field") for h in letter.holes(
        bytes(template.content), "pv_ag")}  # fmt: skip
    return template, fields


def test_the_tool_guesses_the_holes_of_an_ag_model(admin_web):
    template, fields = install(admin_web)
    found = list(fields.values())
    for code in ("client.denomination", "client.adresse", "client.localite", "client.numero",
                 "ag.annee", "ag.date", "ag.debut", "ag.fin", "client.representant",
                 "exercice.cloture", "cabinet.nom", "demander"):  # fmt: skip
        assert code in found, code
    assert None not in found  # chaque trou a une proposition
    labels = [v["label"] for v in template.holes.values() if v["field"] == "demander"]
    assert any("affecter" in label for label in labels)
    # Le modèle de PV n'est pas la lettre de mission : pas de contrôle ITAA, plusieurs actifs.
    page = admin_web.get(reverse("lettre:model", args=[template.pk])).content.decode()
    assert "éléments obligatoires (ITAA)" not in page


def _activate(admin_web, template, dividend_paragraphs=True):
    from apps.documents import letter

    content = bytes(template.content)
    data = {"action": "activer", "forme": ["SRL"]}
    for hole in letter.holes(content, "pv_ag"):
        value = template.holes[hole.key]
        data[f"trou-{hole.key}"] = value["field"]
        data[f"libelle-{hole.key}"] = value["label"]
    if dividend_paragraphs:
        for index, text in letter.paragraphs(content):
            if "reporter" in text:
                data[f"option-{index}"] = "sans_dividende"
            if "dividende brut" in text:
                data[f"option-{index}"] = "dividende"
    admin_web.post(reverse("lettre:model", args=[template.pk]), data)
    template.refresh_from_db()
    assert template.active and template.forms == ["SRL"]


def test_prepare_send_and_sign_the_ag_minutes(admin_web, django_capture_on_commit_callbacks):
    from apps.clients.models import Relation
    from apps.core import events
    from apps.documents import letter
    from apps.documents.models import Document, DocumentRequest

    template, _ = install(admin_web)
    _activate(admin_web, template)
    company = ClientFactory(name="Garage Fictif", legal_form="SRL", vat_status="subject",
                            fiscal_year_end_month=12, fiscal_year_end_day=31,
                            street="Rue Fictive 1", postal_code="5000", city="Namur")  # fmt: skip
    from apps.clients.models import ContactPoint

    ContactPoint.objects.create(client=company, kind=ContactPoint.Kind.EMAIL,
                                value="gerant@exemple.example", is_primary=True)  # fmt: skip
    manager = PersonFactory(name="Gérant", first_name="Paul")
    Relation.objects.create(entity=company, person=manager, role=Relation.Role.MANAGER)
    services.ensure_years(company, date(2026, 10, 30))
    year = ClosingYear.objects.get(client=company)
    # La case AG propose de préparer le PV ; une ASBL sans modèle actif ne l'a pas.
    tab = admin_web.get(reverse("clients:closing", args=[company.pk])).content.decode()
    url = reverse("clients:closing_ag", args=[company.pk, year.pk])
    assert url in tab
    assert letter.ag_templates(ClientFactory(legal_form="ASBL")) == []

    page = admin_web.get(url).content.decode()
    assert "Déjà rempli" in page and "Rue Fictive 1" in page
    questions = letter.questions(template)
    data = {"modele": template.pk, "date": "2026-06-30", "debut": "18:00", "fin": "18:30",
            "dividende": "non", **{k: "1 234,00" for k, _ in questions}}  # fmt: skip
    response = admin_web.post(url, data)
    year.refresh_from_db()
    draft = year.ag_draft
    assert response.status_code == 302 and draft is not None and draft.doc_type == "pv_ag"
    assert draft.state == Document.State.DRAFT and year.dividend is False
    filled = letter.fill(template, company, {k: "1 234,00" for k, _ in questions},
                         ag.context(year, date(2026, 6, 30), *_times(), False, ""))  # fmt: skip
    text = _text(filled)
    assert "Garage Fictif" in text and "30/06/2026" in text and "18h00" in text
    assert "18h30" in text and "31/12/2025" in text and "Paul Gérant" in text
    assert "reporter" in text and "dividende brut" not in text
    # Pas une archive tant que ce n'est qu'un brouillon.
    archives = admin_web.get(reverse("clients:archives", args=[company.pk])).content.decode()
    assert "PV d'AG" not in archives

    # Envoyé pour signature : la case AG est « à signer ».
    admin_web.post(reverse("clients:document_send", args=[company.pk, draft.pk]))
    draft.refresh_from_db()
    assert draft.state == Document.State.AWAITING_SIGNATURE
    tab = admin_web.get(reverse("clients:closing", args=[company.pk])).content.decode()
    assert "à signer" in tab
    # Signé par le client : rangé et la case AG se coche.
    request = DocumentRequest.objects.get(client=company, doc_type="pv_ag")
    signed = Document.objects.create(client=company, title="PV d'AG signé", folder="legal",
                                     doc_type="pv_ag", original_name="pv.pdf", size=10)  # fmt: skip
    with django_capture_on_commit_callbacks(execute=True):
        events.emit("document.signed", subject=company, actor=None,
                    payload={"document": str(signed.pk), "demande": request.reference})  # fmt: skip
    deposit = year.deposits.get(kind=deposits.Kind.AG)
    assert deposit.done_on is not None and list(deposit.documents.all()) == [signed]


def test_dividend_keeps_the_dividend_paragraph_and_adds_the_withholding_box(admin_web):
    from apps.documents import letter

    template, _ = install(admin_web)
    _activate(admin_web, template)
    company = ClientFactory(legal_form="SRL", fiscal_year_end_month=12, fiscal_year_end_day=31)
    services.ensure_years(company, date(2026, 10, 30))
    year = ClosingYear.objects.get(client=company)
    questions = letter.questions(template)
    data = {"modele": template.pk, "date": "2026-06-30", "debut": "18:00", "fin": "18:30",
            "dividende": "oui", **{k: "500,00" for k, _ in questions}}  # fmt: skip
    admin_web.post(reverse("clients:closing_ag", args=[company.pk, year.pk]), data)
    year.refresh_from_db()
    assert year.dividend is True
    text = _text(letter.fill(template, company, {}, ag.context(
        year, date(2026, 6, 30), *_times(), True, "")))  # fmt: skip
    assert "dividende brut" in text and "reporter" not in text
    kinds = [b.kind for b in deposits.boxes(year, deposits.kinds_for(company))]
    assert "pm" in kinds


def test_missing_answers_are_asked_again(admin_web):
    template, _ = install(admin_web)
    _activate(admin_web, template)
    company = ClientFactory(legal_form="SRL", fiscal_year_end_month=12, fiscal_year_end_day=31)
    services.ensure_years(company, date(2026, 10, 30))
    year = ClosingYear.objects.get(client=company)
    url = reverse("clients:closing_ag", args=[company.pk, year.pk])
    response = admin_web.post(url, {"modele": template.pk, "date": "", "debut": "", "fin": ""},
                              follow=True)  # fmt: skip
    assert "À compléter" in response.content.decode()
    year.refresh_from_db()
    assert year.ag_draft is None


def _times():
    from datetime import time

    return time(18), time(18, 30)


def _text(word: bytes) -> str:
    from docx import Document as Word

    return "\n".join(p.text for p in Word(io.BytesIO(word)).paragraphs)


def test_dividend_amount_is_not_required_without_dividend(admin_web):
    from apps.documents import letter

    template, _ = install(admin_web)
    _activate(admin_web, template)
    company = ClientFactory(legal_form="SRL", fiscal_year_end_month=12, fiscal_year_end_day=31)
    services.ensure_years(company, date(2026, 10, 30))
    year = ClosingYear.objects.get(client=company)
    url = reverse("clients:closing_ag", args=[company.pk, year.pk])
    assert 'value="2026-06-30"' in admin_web.get(url).content.decode()  # 6 mois après la clôture
    answers = {k: "100,00" for k, _ in letter.questions(template)
               if letter.condition_of(template, k) != "dividende"}  # fmt: skip
    data = {"modele": template.pk, "date": "2026-06-30", "debut": "18:00", "fin": "18:30",
            "dividende": "non", **answers}  # fmt: skip
    assert admin_web.post(url, data).status_code == 302
    year.refresh_from_db()
    assert year.ag_draft is not None


def test_the_prepared_minutes_can_be_downloaded_in_word(admin_web):
    from apps.documents import letter

    template, _ = install(admin_web)
    _activate(admin_web, template)
    company = ClientFactory(name="Garage Fictif", legal_form="SRL", fiscal_year_end_month=12,
                            fiscal_year_end_day=31)  # fmt: skip
    services.ensure_years(company, date(2026, 10, 30))
    year = ClosingYear.objects.get(client=company)
    word = reverse("clients:closing_ag_word", args=[company.pk, year.pk])
    assert admin_web.get(word).status_code == 404  # rien de préparé
    asked = {k: "1 234,00" for k, _ in letter.questions(template)}
    data = {"modele": template.pk, "date": "2026-06-30", "debut": "18:00", "fin": "18:30",
            "dividende": "non", **asked}  # fmt: skip
    admin_web.post(reverse("clients:closing_ag", args=[company.pk, year.pk]), data)
    tab = admin_web.get(reverse("clients:closing", args=[company.pk])).content.decode()
    assert word in tab
    response = admin_web.get(word)
    assert response.status_code == 200 and "Garage Fictif" in _text(response.content)
    assert "30/06/2026" in _text(response.content)


def model_with_tables() -> bytes:
    from docx import Document as Word

    doc = Word()
    doc.add_paragraph("[SOCIÉTÉ]")
    doc.add_paragraph("La séance est ouverte à [HEURE].")
    doc.add_paragraph("Résolution 1 :")
    doc.add_paragraph("L'assemblée approuve les comptes, bénéfice net déterminé comme suit :")
    first = doc.add_table(rows=0, cols=3)
    for label, sign in (("Bénéfice d'exploitation", "(+)"), ("Charges financières", "(-)"),
                        ("Bénéfice net à affecter", "")):  # fmt: skip
        cells = first.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = label, sign, "[MONTANT]"
    doc.add_paragraph("Résolution 2 :")
    doc.add_paragraph("L'assemblée décide de reporter l'intégralité du bénéfice, soit [MONTANT].")
    second = doc.add_table(rows=0, cols=2)
    for label in ("Rémunération ordinaire", "ATN Auto", "Emoluments de l'exercice"):
        cells = second.add_row().cells
        cells[0].text, cells[1].text = label, "[MONTANT]"
    doc.add_paragraph("Toutes ces résolutions sont adoptées à l'unanimité.")
    doc.add_paragraph("La séance est levée à [HEURE].")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _tables_setup(admin_web):
    from django.core.files.uploadedfile import SimpleUploadedFile

    from apps.documents import letter
    from apps.documents.models import LetterTemplate

    upload = SimpleUploadedFile("pv.docx", model_with_tables())
    admin_web.post(reverse("lettre:models"), {"sorte": "pv_ag", "modele": upload})
    template = LetterTemplate.objects.get()
    data = {"action": "activer", "relu": "1"}
    for hole in letter.holes(bytes(template.content), "pv_ag"):
        value = template.holes[hole.key]
        data[f"trou-{hole.key}"], data[f"libelle-{hole.key}"] = value["field"], value["label"]
    admin_web.post(reverse("lettre:model", args=[template.pk]), data)
    template.refresh_from_db()
    assert template.active
    company = ClientFactory(name="Garage Fictif", legal_form="SRL", fiscal_year_end_month=12,
                            fiscal_year_end_day=31)  # fmt: skip
    services.ensure_years(company, date(2026, 10, 30))
    year = ClosingYear.objects.get(client=company)
    keys = {label: key for key, label in letter.questions(template)}
    return template, company, year, keys


def test_amounts_that_do_not_add_up_are_flagged(admin_web):
    template, company, year, keys = _tables_setup(admin_web)
    assert len(ag.sums(template)) == 2
    amounts = {"Bénéfice d'exploitation": "100000", "Charges financières": "25000",
               "Bénéfice net à affecter": "75000", "Rémunération ordinaire": "20000",
               "ATN Auto": "500", "Emoluments de l'exercice": "5000"}  # fmt: skip
    report = next(label for label in keys if "intégralité" in label)
    data = {"modele": template.pk, "date": "2026-06-30", "debut": "18:00", "fin": "18:30",
            "dividende": "non", keys[report]: "70000",
            **{keys[label]: value for label, value in amounts.items()}}  # fmt: skip
    url = reverse("clients:closing_ag", args=[company.pk, year.pk])
    page = admin_web.post(url, data).content.decode()
    year.refresh_from_db()
    assert year.ag_draft is None  # rien de préparé
    assert "20.500,00" in page and "5.000,00" in page and "75.000,00" in page
    assert "montants_ok" in page
    response = admin_web.post(url, {**data, "montants_ok": "1"})  # confirmé : préparé
    year.refresh_from_db()
    assert response.status_code == 302 and year.ag_draft is not None


def test_correct_amounts_pass(admin_web):
    template, _company, _year, keys = _tables_setup(admin_web)
    report = next(label for label in keys if "intégralité" in label)
    answers = {keys["Bénéfice d'exploitation"]: "100.000,00", keys["Charges financières"]: "25000",
               keys["Bénéfice net à affecter"]: "75000", keys[report]: "75 000",
               keys["Rémunération ordinaire"]: "20000", keys["ATN Auto"]: "500",
               keys["Emoluments de l'exercice"]: "20500"}  # fmt: skip
    assert ag.check_amounts(template, answers) == []


def test_something_particular_becomes_one_more_resolution(admin_web):
    template, company, year, _keys = _tables_setup(admin_web)
    ctx = ag.context(year, date(2026, 6, 30), *_times(), False,
                     "L'assemblée décide d'augmenter la rémunération du gérant.")  # fmt: skip
    text = _text(ag.build(template, company, {}, ctx))
    lines = [line for line in text.splitlines() if line.strip()]
    added = lines.index("Résolution 3 :")
    assert lines[added + 1] == "L'assemblée décide d'augmenter la rémunération du gérant."
    assert lines[added + 2].startswith("Toutes ces résolutions")


def test_a_word_where_an_amount_is_expected_is_refused(admin_web):
    template, company, year, keys = _tables_setup(admin_web)
    data = {"modele": template.pk, "date": "2026-06-30", "debut": "18:00", "fin": "18:30",
            "dividende": "non", **dict.fromkeys(keys.values(), "100"),
            keys["Rémunération ordinaire"]: "Dividende"}  # fmt: skip
    page = admin_web.post(reverse("clients:closing_ag", args=[company.pk, year.pk]), data)
    year.refresh_from_db()
    assert year.ag_draft is None
    assert "« Dividende » n&#x27;est pas un montant" in page.content.decode()
    assert reverse("lettre:model", args=[template.pk]) in page.content.decode()


def test_changing_the_model_changes_the_questions(admin_web):
    from django.core.files.uploadedfile import SimpleUploadedFile

    from apps.documents.models import LetterTemplate

    template, company, year, _keys = _tables_setup(admin_web)
    other = SimpleUploadedFile("autre.docx", model_docx())
    admin_web.post(reverse("lettre:models"), {"sorte": "pv_ag", "modele": other})
    second = LetterTemplate.objects.exclude(pk=template.pk).get()
    _activate(admin_web, second)
    url = reverse("clients:closing_ag", args=[company.pk, year.pk])
    assert "ATN Auto" in admin_web.get(f"{url}?modele={template.pk}").content.decode()
    page = admin_web.get(f"{url}?modele={second.pk}").content.decode()
    assert "ATN Auto" not in page and "data-reload-param" in page


def test_dividend_passages_are_found_at_upload(admin_web):
    from django.core.files.uploadedfile import SimpleUploadedFile

    from apps.documents.models import LetterTemplate

    upload = SimpleUploadedFile("pv.docx", model_with_tables())
    admin_web.post(reverse("lettre:models"), {"sorte": "pv_ag", "modele": upload})
    template = LetterTemplate.objects.get()
    conditions = {v["condition"] for v in template.options.values()}
    assert conditions == {"sans_dividende"}
    page = admin_web.get(reverse("lettre:model", args=[template.pk])).content.decode()
    assert 'value="sans_dividende" selected' in page  # gardé si on enregistre la page


def test_the_assistant_understands_a_particular_request(admin_web, settings, monkeypatch):
    from apps.ai.models import AiCall
    from apps.ai.tests.test_ai import fake_reply
    from apps.documents import letter

    template, company, year, keys = _tables_setup(admin_web)
    report = next(i for i, t in letter.paragraphs(bytes(template.content)) if "intégralité" in t)
    settings.MISTRAL_API_KEY = "fictive"
    seen = []
    fake_reply(monkeypatch, {
        "modifications": [{"paragraphe": report, "action": "remplacer",
                           "texte": "L'assemblée décide de distribuer un dividende brut de "
                                    "50.000,00 € et de reporter le solde, soit 25.000,00 €."}],
        "remarques": ["Le report annoncé (80.000 €) dépasse le bénéfice à affecter."],
    }, seen)  # fmt: skip
    amounts = {"Bénéfice d'exploitation": "100000", "Charges financières": "25000",
               "Bénéfice net à affecter": "75000", "Rémunération ordinaire": "20000",
               "ATN Auto": "500", "Emoluments de l'exercice": "20500"}  # fmt: skip
    data = {"modele": template.pk, "date": "2026-06-30", "debut": "18:00", "fin": "18:30",
            "dividende": "oui", "autres": "distribuer 50 000 de dividende, reporter le reste",
            **{keys[label]: value for label, value in amounts.items()}}  # fmt: skip
    url = reverse("clients:closing_ag", args=[company.pk, year.pk])
    response = admin_web.post(url, data, follow=True)
    assert "dépasse le bénéfice" in response.content.decode()  # remarque montrée
    year.refresh_from_db()
    text = _text(ag.word(year))
    assert "distribuer un dividende brut de 50.000,00 €" in text
    assert "intégralité" not in text and "distribuer 50 000 de dividende" not in text
    call = AiCall.objects.get(feature="pv_particulier")
    assert call.model == "mistral-large-latest"
    assert "Garage Fictif" in seen[0]["messages"][1]["content"]  # le PV rempli est lu


def test_without_the_assistant_the_request_becomes_a_resolution(admin_web, settings):
    settings.MISTRAL_API_KEY = ""
    template, company, _year, _keys = _tables_setup(admin_web)
    edits, remarks = ag.understand(template, company, {}, {"autres": "Une décision."}, None)
    assert edits == [] and remarks == []
