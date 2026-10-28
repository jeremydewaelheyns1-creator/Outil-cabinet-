"""Import de l'export clients FID-Manager sur un classeur fictif."""

import io
from datetime import date

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook

from apps.clients.models import Assignment, Client
from apps.core.models import AccessLog, Portfolio
from apps.core.testing.factories import ClientFactory, UserFactory
from apps.core.testing.numbers import enterprise_number, national_number
from apps.reprise.importer import ImportFileError, parse_fiscal_year_end, run_import
from apps.reprise.models import ImportBatch

pytestmark = pytest.mark.django_db

HEADERS = [
    "Code client", "Nom / Dénomination", "Prénom", "Forme juridique", "N° TVA",
    "N° national", "Rue", "Code postal", "Localité", "Langue", "E-mail", "GSM",
    "Périodicité TVA", "Clôture", "Responsable", "Portefeuille", "Colonne maison",
]  # fmt: skip


def workbook(rows: list[list[object]], *, title_rows: int = 2) -> io.BytesIO:
    wb = Workbook()
    sheet = wb.active
    for _ in range(title_rows):
        sheet.append(["Liste des clients — export fictif"])
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def sample_rows() -> list[list[object]]:
    return [
        ["C001", "Boulangerie Fictive", "", "SRL", f"BE {enterprise_number(1)}", "",
         "Rue du Four 12 bte 3", "5000", "Namur", "Français", "Compta@Fictif.test", "",
         "Trimestriel", "31/12", "AB", "Bureau Namur", "x"],
        ["C002", "Martin", "Julie", "", "", national_number(1)[:-1] and int(national_number(1)),
         "Avenue Fictive", 4000, "Liège", "FR", "", "0470 00 00 00", "", "", "", "", ""],
        ["C003", "Société Invalide", "", "SA", "0123.456.789", "", "", "", "", "NL",
         "", "", "Mensuel", date(2025, 6, 30), "ZZ", "", ""],
        ["C004", "ASBL Les Amis Fictifs", "", "ASBL", "", "", "", "", "", "",
         "", "", "", "", "", "", ""],
        ["C001", "Doublon", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
    ]  # fmt: skip


def test_analysis_changes_nothing_and_reports_everything():
    report = run_import(workbook(sample_rows()), file_name="export.xlsx", apply=False)
    assert Client.objects.count() == 0
    assert report.header_row == 3
    assert report.columns["Nom / Dénomination"] == "name"
    assert report.unmapped == ["Colonne maison"]
    assert report.rows == 5  # la ligne vide est ignorée
    assert report.created == 4
    assert report.skipped == 1
    messages = {(i.row, i.field): i.message for i in report.issues}
    assert "invalide" in messages[(6, "enterprise_number")]
    assert "ligne 4" in messages[(8, "reference")]
    assert report.unknown_initials == ["AB", "ZZ"]
    assert report.portfolios_created == ["Bureau Namur"]
    assert national_number(1) not in str(report.as_json())


def test_import_creates_clean_records():
    UserFactory(initials="AB")
    report = run_import(workbook(sample_rows()), file_name="export.xlsx", apply=True)
    assert report.created == 4
    bakery = Client.objects.get(reference="C001")
    assert bakery.enterprise_number == enterprise_number(1)
    assert (bakery.street, bakery.house_number, bakery.box) == ("Rue du Four", "12", "3")
    assert bakery.vat_periodicity == "quarterly"
    assert (bakery.fiscal_year_end_month, bakery.fiscal_year_end_day) == (12, 31)
    assert bakery.portfolio.name == "Bureau Namur"
    assert bakery.contacts.get().value == "compta@fictif.test"
    assert Assignment.objects.get(client=bakery).role == "manager"

    person = Client.objects.get(reference="C002")
    assert person.kind == Client.Kind.PERSON
    assert person.national_number == national_number(1)
    assert person.birth_date == date(1985, 7, 30)  # déduite du numéro national
    assert person.postal_code == "4000"
    assert person.portfolio == Portfolio.objects.get(name="Cabinet")

    invalid = Client.objects.get(reference="C003")
    assert invalid.enterprise_number is None and invalid.language == "nl"
    assert (invalid.fiscal_year_end_month, invalid.fiscal_year_end_day) == (6, 30)
    assert Client.objects.get(reference="C004").kind == Client.Kind.NONPROFIT


def test_existing_records_are_completed_never_overwritten():
    existing = ClientFactory(
        enterprise_number=enterprise_number(1), name="Nom Actuel", city="", reference=None
    )
    report = run_import(workbook(sample_rows()[:1]), file_name="x.xlsx", apply=True)
    existing.refresh_from_db()
    assert report.created == 0 and report.updated == 1
    assert existing.name == "Nom Actuel"
    assert existing.city == "Namur"
    assert existing.reference == "C001"
    assert [(c.field, c.existing, c.imported) for c in report.conflicts] == [
        ("nom ou dénomination", "Nom Actuel", "Boulangerie Fictive")
    ]


def test_reimporting_the_same_file_changes_nothing():
    run_import(workbook(sample_rows()), file_name="x.xlsx", apply=True)
    report = run_import(workbook(sample_rows()), file_name="x.xlsx", apply=True)
    assert report.created == 0
    assert report.updated == 0
    assert Client.objects.count() == 4


def test_unreadable_or_unrecognized_files():
    with pytest.raises(ImportFileError):
        run_import(io.BytesIO(b"ceci n'est pas un classeur"), file_name="x.xlsx", apply=False)
    wb = Workbook()
    wb.active.append(["Colonne A", "Colonne B"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    with pytest.raises(ImportFileError):
        run_import(buffer, file_name="x.xlsx", apply=False)


@pytest.mark.parametrize(
    ("value", "expected"),
    [("31/12", (12, 31)), ("30-06-2024", (6, 30)), ("décembre", (12, 31)),
     ("février", (2, 28)), (date(2024, 3, 31), (3, 31))],
)  # fmt: skip
def test_fiscal_year_end_formats(value, expected):
    assert parse_fiscal_year_end(value) == expected


def test_web_flow_analysis_then_confirmation(make_user, client):
    importer = make_user(capabilities=["reprise.import"])
    client.force_login(importer)
    upload = SimpleUploadedFile("export.xlsx", workbook(sample_rows()).getvalue())
    response = client.post(reverse("reprise:index"), {"file": upload})
    batch = ImportBatch.objects.get()
    assert response.url == reverse("reprise:detail", args=[batch.pk])
    assert Client.objects.count() == 0
    page = client.get(response.url).content.decode()
    assert "Colonne maison" in page and "Confirmer l'import" in page

    client.post(reverse("reprise:apply", args=[batch.pk]))
    batch.refresh_from_db()
    assert batch.status == ImportBatch.Status.APPLIED
    assert batch.content is None
    assert Client.objects.count() == 4
    assert AccessLog.objects.filter(action=AccessLog.Action.IMPORT).exists()
    # Une seconde confirmation est impossible.
    assert client.post(reverse("reprise:apply", args=[batch.pk])).status_code == 404


def test_reprise_requires_capability(web):
    assert web.get(reverse("reprise:index")).status_code == 403


def test_command(tmp_path):
    from django.core.management import call_command

    path = tmp_path / "export.xlsx"
    path.write_bytes(workbook(sample_rows()).getvalue())
    out = io.StringIO()
    call_command("import_fid_clients", str(path), stdout=out)
    assert "ANALYSE" in out.getvalue() and Client.objects.count() == 0
    call_command("import_fid_clients", str(path), "--appliquer", stdout=out)
    assert Client.objects.count() == 4


def test_bulk_import_defers_bce_completion():
    from apps.core.models import Event

    run_import(workbook(sample_rows()), file_name="x.xlsx", apply=True)
    assert Event.objects.filter(type="client.created", status=Event.Status.PENDING).count() == 4


def test_stale_uploads_are_purged(make_user):
    from datetime import timedelta

    from django.utils import timezone

    from apps.reprise.models import purge_stale_uploads

    owner = make_user()
    old = ImportBatch.objects.create(
        file_name="a.xlsx",
        uploaded_by=owner,
        content=b"x",
        created_at=timezone.now() - timedelta(days=2),
    )
    recent = ImportBatch.objects.create(file_name="b.xlsx", uploaded_by=owner, content=b"y")
    assert purge_stale_uploads() == 1
    old.refresh_from_db()
    recent.refresh_from_db()
    assert old.content is None and old.status == ImportBatch.Status.EXPIRED
    assert recent.content is not None


# --- Dossiers complets : personne de contact, comptable « Hybride », feuille Relations ------


def complete_workbook() -> io.BytesIO:
    wb = Workbook()
    clients = wb.active
    clients.title = "Clients"
    clients.append(
        [
            "Code client",
            "Nom",
            "Prénom",
            "Forme juridique",
            "N° TVA",
            "N° national",
            "Civilité contact",
            "Personne de contact",
            "Expert-comptable",
            "Comptable",
        ]
    )
    clients.append(
        [
            "S1",
            "Holding Fictive",
            "",
            "SRL",
            enterprise_number(11),
            "",
            "Madame",
            "Julie Exemple",
            "AB",
            "Hybride",
        ]
    )
    clients.append(
        ["S2", "Garage Fictif", "", "SRL", enterprise_number(12), "", "Docteur X", "", "AB", "CD"]
    )
    clients.append(["P1", "Exemple", "Paul", "", "", national_number(11), "", "", "AB", ""])
    clients.append(["P2", "Exemple", "Julie", "", "", national_number(12), "", "", "", ""])
    relations = wb.create_sheet("Relations")
    relations.append(["Dossier", "Personne", "Rôle", "Part (%)", "Depuis"])
    relations.append(["S1", "P1", "Gérant", "", "01/01/2020"])
    relations.append(["S1", "P1", "Actionnaire", 0.6, ""])
    relations.append([enterprise_number(12), "S1", "Actionnaire", "100 %", ""])
    relations.append(["P1", "Paul Exemple", "Conjoint", "", ""])  # lui-même
    relations.append(["P1", "P2", "Conjointe", "", ""])
    relations.append(["S2", "Inconnu Fictif", "Gérant", "", ""])
    relations.append(["S2", "P2", "Astronaute", "", ""])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def test_complete_dossiers_with_contact_hybrid_and_relations():
    from apps.clients.models import Relation

    UserFactory(initials="AB")
    UserFactory(initials="CD")
    analysis = run_import(complete_workbook(), file_name="export.xlsx", apply=False)
    assert Relation.objects.count() == 0 and analysis.relations_added == 4
    report = run_import(complete_workbook(), file_name="export.xlsx", apply=True)
    assert (report.created, report.hybrid_set, report.relations_sheet) == (4, 1, "Relations")
    assert (report.relations_rows, report.relations_added) == (7, 4)
    holding = Client.objects.get(reference="S1")
    assert (holding.contact_title, holding.contact_name) == ("Madame", "Julie Exemple")
    assert holding.accountant_hybrid
    assert not holding.assignments.filter(role=Assignment.Role.ASSISTANT).exists()
    assert holding.assignments.get(role=Assignment.Role.MANAGER).user.initials == "AB"
    garage = Client.objects.get(reference="S2")
    assert garage.contact_title == "" and not garage.accountant_hybrid
    paul = Client.objects.get(reference="P1")
    share = Relation.objects.get(entity=holding, person=paul, role="shareholder")
    assert float(share.share_percentage) == 60
    assert Relation.objects.get(entity=holding, person=paul, role="manager").start_date == date(
        2020, 1, 1
    )
    assert float(Relation.objects.get(entity=garage, person=holding).share_percentage) == 100
    assert Relation.objects.filter(entity=paul, role="spouse").count() == 1
    messages = " ".join(i.message for i in report.issues)
    assert "lié à lui-même" in messages and "introuvable" in messages
    assert "Astronaute" in messages and "civilité" in messages
    again = run_import(complete_workbook(), file_name="export.xlsx", apply=True)
    assert again.relations_added == 0 and Relation.objects.count() == 4
    assert national_number(11) not in str(again.as_json())


def test_web_report_shows_relations_and_hybrid(make_user, client):
    from apps.clients.models import Relation

    importer = make_user(capabilities=["reprise.import"], initials="AB")
    client.force_login(importer)
    upload = SimpleUploadedFile("export.xlsx", complete_workbook().getvalue())
    response = client.post(reverse("reprise:index"), {"file": upload})
    page = client.get(response.url).content.decode()
    assert "Relations" in page and "4 liens à créer" in page and "Comptable « Hybride »" in page
    client.post(reverse("reprise:apply", args=[ImportBatch.objects.get().pk]))
    assert Relation.objects.count() == 4


@pytest.mark.parametrize(
    ("value", "expected"),
    [("Belge", "BE"), ("Belgique", "BE"), ("Française", "FR"), ("BE", "BE"), ("Martienne", None)],
)
def test_nationalities_as_written_in_fid(value, expected):
    from apps.reprise.importer import _country

    assert _country(value) == expected
