"""Import de l'export « clients » de FID-Manager (fichier Excel).

Déroulement :
1. analyse : le fichier est importé dans une transaction annulée à la fin. On obtient le
   rapport exact de ce qui se passera (créations, mises à jour, erreurs) sans rien modifier ;
2. import : le même code, cette fois validé.

Règles de prudence :
- une fiche existante (même numéro d'entreprise, numéro national ou référence) n'est jamais
  écrasée : seuls ses champs vides sont complétés ; les valeurs différentes sont signalées ;
- un numéro invalide est ignoré et signalé, la ligne est importée sans lui ;
- le rapport ne contient jamais de numéro national.
"""

from __future__ import annotations

import calendar
import contextlib
import functools
import operator
import re
from collections.abc import Iterator
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from functools import cache
from pathlib import Path
from typing import IO, Any

import yaml
from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q
from openpyxl import load_workbook
from stdnum.be import nn

from apps.clients.identifiers import normalize_enterprise_number, normalize_national_number
from apps.clients.models import Assignment, Client, ContactPoint, Relation, field_label
from apps.core import events
from apps.core.models import Portfolio, User
from apps.core.text import normalize_key, normalize_text

MAPPING_FILE = Path(settings.BASE_DIR) / "config" / "reprise" / "clients.yaml"
HEADER_SCAN_ROWS = 15
SOURCE = "reprise"
DEFAULT_PORTFOLIO = "Cabinet"

CLIENT_FIELDS = (
    "reference", "name", "first_name", "legal_form", "kind", "status", "enterprise_number",
    "national_number", "birth_date", "nationality", "language", "street", "house_number", "box",
    "postal_code", "city", "country", "vat_status", "vat_periodicity", "client_since",
    "client_until", "notes", "contact_name", "contact_title", "accountant_hybrid",
)  # fmt: skip
HYBRID = "hybride"  # valeur de la colonne du comptable : gestion hybride (D-063)
SENSITIVE_FIELDS = {"national_number"}
CONTACT_FIELDS: dict[str, str] = {
    "email": ContactPoint.Kind.EMAIL,
    "phone": ContactPoint.Kind.PHONE,
    "mobile": ContactPoint.Kind.PHONE,
    "website": ContactPoint.Kind.WEBSITE,
}
ASSIGNMENT_FIELDS: dict[str, str] = {
    "partner": Assignment.Role.PARTNER,
    "manager": Assignment.Role.MANAGER,
    "assistant": Assignment.Role.ASSISTANT,
}
MONTHS = {
    normalize_key(name): number
    for number, names in enumerate(
        [
            [], ["janvier", "jan"], ["fevrier", "fev"], ["mars"], ["avril", "avr"], ["mai"],
            ["juin"], ["juillet", "juil"], ["aout"], ["septembre", "sept"],
            ["octobre", "oct"], ["novembre", "nov"], ["decembre", "dec"],
        ]
    )
    for name in names
}  # fmt: skip
COUNTRIES = {
    "belgique": "BE", "belgie": "BE", "belgium": "BE", "france": "FR", "pays bas": "NL",
    "nederland": "NL", "luxembourg": "LU", "allemagne": "DE", "deutschland": "DE",
    "espagne": "ES", "italie": "IT", "suisse": "CH", "royaume uni": "GB", "portugal": "PT",
    # Nationalités telles qu'écrites dans FID-Manager (« Belge », « Française »…).
    "belge": "BE", "francaise": "FR", "francais": "FR", "neerlandaise": "NL",
    "neerlandais": "NL", "luxembourgeoise": "LU", "luxembourgeois": "LU", "allemande": "DE",
    "allemand": "DE", "espagnole": "ES", "espagnol": "ES", "italienne": "IT", "italien": "IT",
    "britannique": "GB", "portugaise": "PT", "portugais": "PT",
}  # fmt: skip
STREET_WITH_NUMBER = re.compile(
    r"^(?P<street>.*?\D)[\s,]+(?P<number>\d+[A-Za-z]?)"
    r"(?:\s*[,/]?\s*(?:bte|bt|bus|boite|boîte|b)\.?\s*(?P<box>\S+))?$",
    re.IGNORECASE,
)


class ImportFileError(Exception):
    """Fichier illisible ou sans en-tête reconnaissable."""


class _DryRun(Exception):
    """Annule la transaction à la fin d'une analyse."""


# --- Paramétrage ----------------------------------------------------------------------------


@dataclass(frozen=True)
class Mapping:
    aliases: dict[str, str]  # en-tête normalisé → champ
    values: dict[str, dict[str, str]]  # champ → valeur normalisée → code
    nonprofit_forms: frozenset[str]
    relation_aliases: dict[str, str] = field(default_factory=dict)  # feuille « Relations »
    relation_roles: dict[str, str] = field(default_factory=dict)  # valeur normalisée → rôle


@cache
def load_mapping(path: Path = MAPPING_FILE) -> Mapping:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    aliases: dict[str, str] = {}
    for target, headers in data["champs"].items():
        for header in headers:
            key = normalize_key(header)
            if key in aliases and aliases[key] != target:
                raise ImportFileError(
                    f"Intitulé « {header} » attribué à deux champs : {aliases[key]} et {target}."
                )
            aliases[key] = target
    values = {
        target: {normalize_key(alias): code for code, alias_list in choices.items()
                 for alias in alias_list}
        for target, choices in data.get("valeurs", {}).items()
    }  # fmt: skip
    forms = frozenset(normalize_key(f) for f in data.get("formes_non_marchandes", []))
    relations = data.get("relations", {})
    relation_aliases = {
        normalize_key(header): target
        for target, headers in relations.get("champs", {}).items()
        for header in headers
    }
    relation_roles = {
        normalize_key(alias): code
        for code, alias_list in relations.get("roles", {}).items()
        for alias in alias_list
    }
    return Mapping(aliases, values, forms, relation_aliases, relation_roles)


# --- Rapport --------------------------------------------------------------------------------


@dataclass
class Issue:
    row: int
    field: str
    message: str
    level: str = "error"  # "error" : ligne non importée ; "warning" : importée avec réserve


@dataclass
class Conflict:
    row: int
    client: str
    field: str
    existing: str
    imported: str


@dataclass
class Report:
    file_name: str
    sheet: str = ""
    header_row: int = 0
    columns: dict[str, str] = field(default_factory=dict)  # en-tête → champ
    unmapped: list[str] = field(default_factory=list)
    duplicated_columns: list[str] = field(default_factory=list)
    rows: int = 0
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    skipped: int = 0
    contacts_added: int = 0
    assignments_added: int = 0
    portfolios_created: list[str] = field(default_factory=list)
    unknown_initials: list[str] = field(default_factory=list)
    hybrid_set: int = 0
    relations_sheet: str = ""
    relations_rows: int = 0
    relations_added: int = 0
    issues: list[Issue] = field(default_factory=list)
    conflicts: list[Conflict] = field(default_factory=list)
    applied: bool = False

    def as_json(self) -> dict[str, Any]:
        return asdict(self)

    @property
    def errors(self) -> list[Issue]:
        return [i for i in self.issues if i.level == "error"]

    @property
    def warnings(self) -> list[Issue]:
        return [i for i in self.issues if i.level == "warning"]


# --- Lecture du fichier ---------------------------------------------------------------------


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_rows(
    source: Path | IO[bytes], mapping: Mapping, report: Report
) -> Iterator[tuple[int, dict[str, Any]]]:
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError("Fichier illisible : un classeur Excel (.xlsx) est attendu.") from exc
    try:
        sheet = workbook.worksheets[0]
        report.sheet = sheet.title
        rows = sheet.iter_rows(values_only=True)
        buffered: list[tuple[Any, ...]] = []
        best_index, best_score = -1, 1
        for index, row in enumerate(rows):
            buffered.append(row)
            score = sum(1 for c in row if normalize_key(_cell_text(c)) in mapping.aliases)
            if score > best_score:
                best_index, best_score = index, score
            if index + 1 >= HEADER_SCAN_ROWS:
                break
        if best_index < 0:
            raise ImportFileError(
                "Aucune ligne d'en-tête reconnue. Ajoutez les intitulés de colonnes dans "
                "config/reprise/clients.yaml."
            )
        report.header_row = best_index + 1
        positions: dict[int, str] = {}
        for position, cell in enumerate(buffered[best_index]):
            header = _cell_text(cell)
            if not header:
                continue
            target = mapping.aliases.get(normalize_key(header))
            if target is None:
                report.unmapped.append(header)
            elif target in positions.values():
                report.duplicated_columns.append(header)
            else:
                positions[position] = target
                report.columns[header] = target

        def records(
            source_rows: Iterator[tuple[Any, ...]], start: int
        ) -> Iterator[tuple[int, dict[str, Any]]]:
            for offset, row in enumerate(source_rows):
                values = {
                    target: row[position]
                    for position, target in positions.items()
                    if position < len(row) and _cell_text(row[position])
                }
                if values:
                    yield start + offset, values

        first_data = best_index + 2
        yield from records(iter(buffered[best_index + 1 :]), first_data)
        yield from records(rows, len(buffered) + 1)
    finally:
        workbook.close()


# --- Conversion des valeurs -----------------------------------------------------------------


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _cell_text(value)
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"date non reconnue : {text}")


def parse_fiscal_year_end(value: Any) -> tuple[int, int]:
    """'31/12', '31-12-2025', une date Excel ou 'décembre' → (mois, jour)."""
    if isinstance(value, datetime | date):
        return value.month, value.day
    text = _cell_text(value)
    month = MONTHS.get(normalize_key(text))
    if month:
        return month, calendar.monthrange(2023, month)[1]
    match = re.fullmatch(r"(\d{1,2})\s*[/.\-]\s*(\d{1,2})(?:\s*[/.\-]\s*\d{2,4})?", text)
    if match:
        day, month = int(match.group(1)), int(match.group(2))
        date(2024, month, day)  # lève ValueError si la date n'existe pas
        return month, day
    raise ValueError(f"date de clôture non reconnue : {text}")


def _choice(mapping: Mapping, target: str, value: Any) -> str | None:
    return mapping.values.get(target, {}).get(normalize_key(_cell_text(value)))


def _country(value: Any) -> str | None:
    text = _cell_text(value)
    if re.fullmatch(r"[A-Za-z]{2}", text):
        return text.upper()
    return COUNTRIES.get(normalize_key(text))


@dataclass
class ParsedRow:
    values: dict[str, Any]
    contacts: list[tuple[str, str]]
    assignments: dict[str, str]
    portfolio: str


def parse_row(row_number: int, raw: dict[str, Any], mapping: Mapping, report: Report) -> ParsedRow:
    values: dict[str, Any] = {}

    def warn(field_name: str, message: str) -> None:
        report.issues.append(Issue(row_number, field_name, message, "warning"))

    for name in ("reference", "name", "first_name", "legal_form", "nationality", "street",
                 "house_number", "box", "postal_code", "city", "notes",
                 "contact_name"):  # fmt: skip
        if name in raw:
            values[name] = _cell_text(raw[name])
    if "nationality" in values:
        nationality = _country(values["nationality"])
        values["nationality"] = nationality or ""
        if nationality is None:
            warn("nationality", "nationalité non reconnue, ignorée")

    if "enterprise_number" in raw:
        try:
            values["enterprise_number"] = (
                normalize_enterprise_number(_cell_text(raw["enterprise_number"])) or None
            )
        except ValidationError:
            warn("enterprise_number", "numéro d'entreprise invalide, ignoré")
    if "national_number" in raw:
        text = _cell_text(raw["national_number"])
        if text.isdigit() and len(text) == 10:
            text = text.zfill(11)  # zéro initial perdu par Excel
        try:
            values["national_number"] = normalize_national_number(text) or None
        except ValidationError:
            warn("national_number", "numéro national invalide, ignoré")

    for name in ("birth_date", "client_since", "client_until"):
        if name in raw:
            try:
                values[name] = parse_date(raw[name])
            except ValueError:
                warn(name, "date non reconnue, ignorée")
    if not values.get("birth_date") and values.get("national_number"):
        # Numéro « bis » : la date de naissance peut être inconnue.
        with contextlib.suppress(Exception):
            values["birth_date"] = nn.get_birth_date(values["national_number"])

    if "fiscal_year_end" in raw:
        try:
            values["fiscal_year_end_month"], values["fiscal_year_end_day"] = parse_fiscal_year_end(
                raw["fiscal_year_end"]
            )
        except ValueError:
            warn("fiscal_year_end", "date de clôture non reconnue, ignorée")

    for name in ("language", "status", "vat_status", "vat_periodicity", "kind"):
        if name in raw:
            code = _choice(mapping, name, raw[name])
            if code is None:
                warn(name, f"valeur « {_cell_text(raw[name])} » non reconnue, ignorée")
            else:
                values[name] = code

    if "contact_title" in raw:
        titles = {normalize_key(label): code for code, label in Client.Title.choices if code}
        title = titles.get(normalize_key(_cell_text(raw["contact_title"])))
        if title is None:
            warn("contact_title", "civilité non reconnue, ignorée")
        else:
            values["contact_title"] = title

    if "country" in raw:
        country = _country(raw["country"])
        if country is None:
            warn("country", f"pays « {_cell_text(raw['country'])} » non reconnu, ignoré")
        else:
            values["country"] = country

    # « Rue de la Station 12 bte 3 » dans une seule colonne : on sépare le numéro et la boîte.
    if values.get("street") and not values.get("house_number"):
        match = STREET_WITH_NUMBER.match(values["street"])
        if match:
            values["street"] = match.group("street").strip(" ,")
            values["house_number"] = match.group("number")
            if match.group("box"):
                values["box"] = match.group("box")

    if "kind" not in values:
        form = normalize_key(values.get("legal_form", ""))
        if form in mapping.nonprofit_forms:
            values["kind"] = Client.Kind.NONPROFIT
        elif values.get("national_number") or (values.get("first_name") and not form):
            values["kind"] = Client.Kind.PERSON
        else:
            values["kind"] = Client.Kind.COMPANY

    contacts = [
        (CONTACT_FIELDS[name], _cell_text(raw[name]).lower() if name == "email"
         else _cell_text(raw[name]))
        for name in CONTACT_FIELDS
        if name in raw
    ]  # fmt: skip
    assignments = {
        role: _cell_text(raw[name]) for name, role in ASSIGNMENT_FIELDS.items() if name in raw
    }
    if normalize_key(assignments.get(Assignment.Role.ASSISTANT, "")) == HYBRID:
        assignments.pop(Assignment.Role.ASSISTANT)
        values["accountant_hybrid"] = True
    return ParsedRow(values, contacts, assignments, _cell_text(raw.get("portfolio")))


# --- Import ---------------------------------------------------------------------------------


class _Importer:
    def __init__(self, report: Report, actor: User | None) -> None:
        self.report = report
        self.actor = actor
        self.portfolios = {normalize_text(p.name): p for p in Portfolio.objects.all()}
        self.users: dict[str, User] = {}
        for user in User.objects.filter(is_active=True):
            for key in (user.initials, user.email, user.get_full_name()):
                if key:
                    self.users.setdefault(normalize_text(key), user)
        self.seen: dict[tuple[str, str], int] = {}

    def portfolio(self, name: str) -> Portfolio:
        name = name.strip() or DEFAULT_PORTFOLIO
        key = normalize_text(name)
        if key not in self.portfolios:
            self.portfolios[key] = Portfolio.objects.create(name=name[:100])
            self.report.portfolios_created.append(name[:100])
        return self.portfolios[key]

    def find_existing(self, row_number: int, values: dict[str, Any]) -> Client | bool | None:
        """Fiche existante, None si nouvelle, False si la ligne est ambiguë."""
        matches: dict[int, Client] = {}
        for key in ("enterprise_number", "national_number", "reference"):
            value = values.get(key)
            if not value:
                continue
            previous_row = self.seen.get((key, value))
            if previous_row is not None:
                self.report.issues.append(
                    Issue(row_number, key, f"même identifiant qu'à la ligne {previous_row}")
                )
                return False
            client = Client.objects.filter(**{key: value}).first()
            if client is not None:
                matches[client.pk] = client
        if len(matches) > 1:
            self.report.issues.append(
                Issue(row_number, "", "les identifiants désignent des fiches différentes")
            )
            return False
        found = next(iter(matches.values()), None)
        if found is not None and found.deleted_at is not None:
            # D-095 : jamais de mise à jour discrète d'un dossier mis à la corbeille.
            self.report.issues.append(
                Issue(row_number, "", "ce dossier est dans la corbeille : réactivez-le d'abord")
            )
            return False
        return found

    def remember(self, row_number: int, values: dict[str, Any]) -> None:
        for key in ("enterprise_number", "national_number", "reference"):
            if values.get(key):
                self.seen[(key, values[key])] = row_number

    def import_row(self, row_number: int, parsed: ParsedRow) -> None:
        values = parsed.values
        if not values.get("name"):
            self.report.issues.append(Issue(row_number, "name", "nom manquant"))
            self.report.skipped += 1
            return
        existing = self.find_existing(row_number, values)
        if existing is False:
            self.report.skipped += 1
            return
        self.remember(row_number, values)
        if isinstance(existing, Client):
            client = existing
            changed = self.complete(row_number, client, values)
        else:
            client = Client(portfolio=self.portfolio(parsed.portfolio))
            for name, value in values.items():
                setattr(client, name, value)
            changed = True
        if changed:
            try:
                with transaction.atomic():
                    client.full_clean()
                    client.save()
            except ValidationError as exc:
                for field_name, messages in exc.message_dict.items():
                    for message in messages:
                        self.report.issues.append(Issue(row_number, field_name, message))
                self.report.skipped += 1
                return
        if existing is None:
            self.report.created += 1
            events.emit(
                "client.created",
                subject=client,
                actor=self.actor,
                payload={"source": SOURCE},
                process_now=False,  # complétion BCE en arrière-plan : import de masse
            )
        elif changed:
            self.report.updated += 1
        else:
            self.report.unchanged += 1
        self.add_contacts(client, parsed.contacts)
        self.add_assignments(client, parsed.assignments)
        if values.get("accountant_hybrid"):
            self.report.hybrid_set += 1

    def complete(self, row_number: int, client: Client, values: dict[str, Any]) -> bool:
        changed = False
        for name, value in values.items():
            if value in (None, ""):
                continue
            current = getattr(client, name)
            if isinstance(current, bool):  # « Hybride » : seulement sans comptable fixe
                fixed = client.assignments.filter(role=Assignment.Role.ASSISTANT).exists()
                if value and not current and not fixed:
                    setattr(client, name, value)
                    changed = True
                continue
            if current in (None, ""):
                setattr(client, name, value)
                changed = True
            elif normalize_text(str(current)) != normalize_text(str(value)):
                hidden = name in SENSITIVE_FIELDS
                self.report.conflicts.append(
                    Conflict(
                        row_number,
                        client.display_name,
                        field_label(name),
                        "•••" if hidden else str(current),
                        "•••" if hidden else str(value),
                    )
                )
        return changed

    def add_contacts(self, client: Client, contacts: list[tuple[str, str]]) -> None:
        for kind, value in contacts:
            if not value:
                continue
            _, created = ContactPoint.objects.get_or_create(
                client=client, kind=kind, value=value[:254], defaults={"source": SOURCE}
            )
            self.report.contacts_added += int(created)

    def add_assignments(self, client: Client, assignments: dict[str, str]) -> None:
        for role, key in assignments.items():
            if not key:
                continue
            user = self.users.get(normalize_text(key))
            if user is None:
                if key not in self.report.unknown_initials:
                    self.report.unknown_initials.append(key)
                continue
            _, created = Assignment.objects.get_or_create(
                client=client, role=role, defaults={"user": user}
            )
            self.report.assignments_added += int(created)


# --- Feuille « Relations » -------------------------------------------------------------------


def _names(client: Client) -> set[str]:
    first, name = client.first_name, client.name
    return {normalize_text(n) for n in (client.display_name, name, f"{first} {name}",
                                        f"{name} {first}")}  # fmt: skip


def _find_client(text: str) -> Client | None:
    """Une fiche désignée par son code client, son numéro d'entreprise ou national, ou son nom
    exact (s'il est unique)."""
    text = text.strip()
    if not text:
        return None
    found = Client.objects.filter(reference__iexact=text).first()
    if found is not None:
        return found
    for normalize, field_name in ((normalize_enterprise_number, "enterprise_number"),
                                  (normalize_national_number, "national_number")):  # fmt: skip
        with contextlib.suppress(ValidationError):
            number = normalize(text)
            if number:
                found = Client.objects.filter(**{field_name: number}).first()
                if found is not None:
                    return found
    wanted = normalize_text(text)
    candidates = Client.objects.filter(
        functools.reduce(operator.or_, (Q(name__icontains=w) for w in text.split()))
    )
    by_name = [c for c in candidates if wanted in _names(c)]
    return by_name[0] if len(by_name) == 1 else None


def import_relations(source: Path | IO[bytes], mapping: Mapping, report: Report) -> None:
    """Feuille « Relations » (facultative) : gérants, administrateurs, actionnaires, conjoints.
    Un lien déjà présent (même dossier, même personne, même rôle) n'est jamais doublé."""
    if not mapping.relation_aliases:
        return
    if hasattr(source, "seek"):
        source.seek(0)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets[1:]:
            rows = list(sheet.iter_rows(values_only=True))
            header_index = next((i for i, row in enumerate(rows[:HEADER_SCAN_ROWS])
                                 if sum(1 for c in row if normalize_key(_cell_text(c))
                                        in mapping.relation_aliases) >= 3), None)  # fmt: skip
            if header_index is None:
                continue
            report.relations_sheet = sheet.title
            positions = {
                i: mapping.relation_aliases[normalize_key(_cell_text(c))]
                for i, c in enumerate(rows[header_index])
                if normalize_key(_cell_text(c)) in mapping.relation_aliases
            }
            for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
                raw = {
                    t: row[i] for i, t in positions.items() if i < len(row) and _cell_text(row[i])
                }
                if not raw:
                    continue
                report.relations_rows += 1
                _import_relation(offset, raw, mapping, report)
            return
    finally:
        workbook.close()


def _import_relation(row: int, raw: dict[str, Any], mapping: Mapping, report: Report) -> None:
    def error(message: str) -> None:
        report.issues.append(Issue(row, "relations", f"Relations : {message}"))

    entity = _find_client(_cell_text(raw.get("entity")))
    person = _find_client(_cell_text(raw.get("person")))
    role = mapping.relation_roles.get(normalize_key(_cell_text(raw.get("role"))))
    if entity is None or person is None:
        error("dossier ou personne introuvable (code client, numéro ou nom exact)")
        return
    if entity.pk == person.pk:
        error("un dossier ne peut pas être lié à lui-même")
        return
    if role is None:
        error(f"rôle « {_cell_text(raw.get('role'))} » non reconnu")
        return
    share = None
    if raw.get("share") not in (None, ""):
        text = _cell_text(raw["share"]).replace("%", "").replace(",", ".").strip()
        try:
            share = float(text)
            if 0 < share <= 1 and isinstance(raw["share"], int | float):
                share *= 100  # cellule Excel au format pourcentage (0,8 → 80 % ; 1 → 100 %)
        except ValueError:
            report.issues.append(
                Issue(row, "relations", "Relations : pourcentage non reconnu, ignoré", "warning")
            )
    dates = {}
    for name in ("start", "end"):
        if raw.get(name):
            with contextlib.suppress(ValueError):
                dates[f"{name}_date"] = parse_date(raw[name])
    if Relation.objects.filter(entity=entity, person=person, role=role).exists():
        return
    relation = Relation(entity=entity, person=person, role=role, source=SOURCE, **dates)
    if share is not None:
        relation.share_percentage = round(share, 2)
    try:
        with transaction.atomic():
            relation.full_clean()
            relation.save()
    except ValidationError as exc:
        error(" ".join(m for ms in exc.message_dict.values() for m in ms))
        return
    report.relations_added += 1


def run_import(
    source: Path | IO[bytes], *, file_name: str, apply: bool, actor: User | None = None
) -> Report:
    """Analyse (apply=False) ou importe (apply=True) un export clients FID-Manager."""
    mapping = load_mapping()
    report = Report(file_name=file_name)
    try:
        with transaction.atomic():
            importer = _Importer(report, actor)
            for row_number, raw in _read_rows(source, mapping, report):
                report.rows += 1
                importer.import_row(row_number, parse_row(row_number, raw, mapping, report))
            import_relations(source, mapping, report)
            if not apply:
                raise _DryRun
    except _DryRun:
        pass
    report.applied = apply
    if "name" not in report.columns.values():
        report.issues.insert(0, Issue(report.header_row, "name", "aucune colonne « Nom » reconnue"))
    return report
